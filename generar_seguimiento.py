# -*- coding: utf-8 -*-
"""
Genera demo-interno/data/seguimiento.json (contrato `seguimiento` v2, ver demo-interno/SCHEMA.md) a partir de la
hoja `evol_conc` de Res_firmadas.xlsm: las concentraciones EN TRAMITE que ve el area interna.

Que es "en tramite": la misma definicion del Excel (hoja General, celda C26 "TOTAL ACTIVAS"): ESTADO en
Pendiente de presentacion, Observado, En analisis, En instruccion, Suspendida, IT circulando, Para resolucion o TDC.
El script compara su conteo con el que el Excel tiene guardado y avisa si no coinciden.

Que muestra de cada expediente:
  - id           CONC-NNNN (columna Clave)
  - operacion    la caratula RESUMIDA: sin comillas ni "S/ NOTIFICACION ART. 9 DE LA LEY 27.442" ni "(CONC NNNN)"
  - abogados     Abogado_1 y Abogado_2, con el apellido normalizado
  - economistas  Economista_1 y Economista_2
  - tipo         ES FT = SI -> PROSUM, NO -> Ordinario
  - estado, ingreso (Fecha_Ingreso)

El Excel se abre en SOLO LECTURA (las macros no se ejecutan). El unico archivo que se escribe es el JSON.
El JSON trae expedientes en tramite: es de uso interno (ver SITIO.md, "Area interna").

Uso:
    python -B generar_seguimiento.py
    python -B generar_seguimiento.py --excel RUTA/Res_firmadas.xlsm --salida RUTA/seguimiento.json --corte 2026-10-09
Necesita openpyxl (el resto del circuito ya lo usa).
"""
from __future__ import annotations

import argparse
import collections
import datetime as dt
import json
import re
import sys
import unicodedata
from pathlib import Path

AQUI = Path(__file__).resolve().parent
EXCEL = AQUI / "Res_firmadas.xlsm"
SALIDA = AQUI / "demo-interno" / "data" / "seguimiento.json"

# Estados de la hoja evol_conc que el Excel cuenta como "en tramite" (General!C12:C19), en el orden del tramite
# (el mismo de la lista desplegable de la hoja oculta _listas).
ESTADOS_ACTIVOS = ["Pendiente de presentación", "Observado", "En análisis", "En instrucción",
                   "Suspendida", "IT circulando", "Para resolución", "TDC"]
HOJA = "evol_conc"
COLUMNAS = ["ESTADO", "Clave", "Carpeta", "Caratula", "Fecha_Ingreso", "Fecha_firma",
            "Abogado_1", "Abogado_2", "Economista_1", "Economista_2", "ES FT"]
TIPO_POR_ES_FT = {"SI": "PROSUM", "NO": "Ordinario"}
MAX_CARATULA = 120

# Apellidos mal tipeados en el Excel -> el apellido correcto (ambos sin tildes y en mayusculas).
ALIAS_APELLIDOS = {"ROSOZKA": "ROSOSZKA", "ROZOSKA": "ROSOSZKA"}

AVISO_INTERNO = ("USO INTERNO: expedientes en trámite tomados de Res_firmadas.xlsm (hoja evol_conc). "
                 "No difundir fuera del área.")


# --------------------------------------------------------------------------- #
# Texto
# --------------------------------------------------------------------------- #
def limpiar(txt) -> str:
    """Texto sin espacios duros ni repetidos."""
    if txt is None:
        return ""
    return re.sub(r"\s+", " ", str(txt).replace("\xa0", " ")).strip()


RX_COMILLAS = re.compile(r'["“”«»]')
RX_PREFIJO = re.compile(r"^\s*CONC(?:\.|\s)*(?:N\s*[°º]\s*)?\d+\s*[-–—:]+\s*", re.I)
RX_NOTIF = re.compile(r"\bS\s*/\s*NOTIF\w*\.?", re.I)
RX_CONC = re.compile(r"\(\s*CONC[^)]*\)", re.I)
RX_PARENTESIS = re.compile(r"\(([^()]*)\)")


def resumir_caratula(texto, max_len: int = MAX_CARATULA) -> str:
    """'ABBOTT LABORATORIES y POSITRON LIMITED S/ NOTIFICACION ART. 8 DE LA LEY 25.156 (CONC. N° 1169)'
    -> 'ABBOTT LABORATORIES y POSITRON LIMITED'. Se queda con las partes; de lo que viene despues de
    "S/ NOTIFICACION" solo conserva un paréntesis que aclare una parte (no el '(CONC nnnn)')."""
    original = limpiar(texto)
    t = limpiar(RX_COMILLAS.sub("", original))
    t = RX_PREFIJO.sub("", t, count=1)
    m = RX_NOTIF.search(t)
    partes, resto = (t[:m.start()], t[m.end():]) if m else (t, "")
    partes = RX_CONC.sub("", partes)
    extras = [limpiar(p) for p in RX_PARENTESIS.findall(RX_CONC.sub("", resto)) if limpiar(p)]
    res = limpiar(partes).strip(" ,;-–—:")
    for e in extras:
        res += " (%s)" % e
    res = res or original
    if len(res) > max_len:
        res = res[:max_len - 1].rsplit(" ", 1)[0].rstrip(" ,;-–—:(") + "…"
    return res


def plegar(s) -> str:
    """Apellido en mayusculas, sin tildes ni espacios repetidos: la clave para juntar variantes."""
    base = unicodedata.normalize("NFD", limpiar(s))
    return "".join(c for c in base if not unicodedata.combining(c)).upper()


def con_tilde(s: str) -> bool:
    return plegar(s) != limpiar(s).upper()


def nombres_unicos(filas, columnas):
    """Ordena y unifica los apellidos de las columnas dadas. Devuelve {clave: nombre_a_mostrar}.
    ZUVIRIA / ZUVIRÍA -> 'Zuviría'; ROSOZKA / Rososzka -> 'Rososzka' (por ALIAS_APELLIDOS)."""
    variantes = collections.defaultdict(collections.Counter)
    for fila in filas:
        for col in columnas:
            v = limpiar(fila.get(col))
            if v:
                variantes[ALIAS_APELLIDOS.get(plegar(v), plegar(v))][v] += 1
    mostrar = {}
    for clave, cuenta in variantes.items():
        propias = {v: n for v, n in cuenta.items() if plegar(v) == clave} or cuenta
        elegido = max(propias, key=lambda v: (con_tilde(v), propias[v]))
        mostrar[clave] = elegido.title()
    return mostrar


def equipo(fila, columnas, mostrar):
    out = []
    for col in columnas:
        v = limpiar(fila.get(col))
        if v:
            nombre = mostrar[ALIAS_APELLIDOS.get(plegar(v), plegar(v))]
            if nombre not in out:
                out.append(nombre)
    return out


def a_fecha(v):
    """datetime/date/texto -> date; None si no es una fecha valida."""
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = limpiar(v)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def id_de(fila) -> str:
    clave = limpiar(fila.get("Clave"))
    if not clave:
        carpeta = limpiar(fila.get("Carpeta"))
        clave = "CONC %s" % carpeta if carpeta else ""
    return clave.replace(" ", "-")


# --------------------------------------------------------------------------- #
# Excel (solo lectura)
# --------------------------------------------------------------------------- #
def leer_excel(ruta: Path):
    """-> (filas de evol_conc como dicts, total de activas que guardo el Excel, 'fecha de corte' del Excel)."""
    try:
        import openpyxl
    except ImportError:
        raise SystemExit("Falta openpyxl: pip install openpyxl")
    if not ruta.exists():
        raise SystemExit("No encuentro %s (pasar --excel RUTA)." % ruta)
    wb = openpyxl.load_workbook(str(ruta), read_only=True, data_only=True)
    if HOJA not in wb.sheetnames:
        raise SystemExit("%s no tiene la hoja '%s' (hojas: %s)." % (ruta.name, HOJA, ", ".join(wb.sheetnames)))
    it = wb[HOJA].iter_rows(values_only=True)
    enc = [limpiar(c) for c in next(it)]
    faltan = [c for c in COLUMNAS if c not in enc]
    if faltan:
        raise SystemExit("La hoja %s no tiene las columnas: %s." % (HOJA, ", ".join(faltan)))
    filas = [dict(zip(enc, r)) for r in it if any(v not in (None, "") for v in r)]

    total_excel = corte_excel = None
    if "General" in wb.sheetnames:
        g = wb["General"]
        try:
            total_excel = g["C26"].value
            corte_excel = a_fecha(g["C5"].value)
        except (KeyError, IndexError, TypeError):
            pass
    wb.close()
    return filas, total_excel, corte_excel


# --------------------------------------------------------------------------- #
# Armado del JSON
# --------------------------------------------------------------------------- #
def construir(filas, corte: dt.date, total_excel=None, corte_excel=None, fuente="Res_firmadas.xlsm › evol_conc"):
    avisos = []
    activas = [f for f in filas if limpiar(f.get("ESTADO")) in ESTADOS_ACTIVOS]
    sin_estado = [id_de(f) for f in filas if not limpiar(f.get("ESTADO"))]
    if sin_estado:
        avisos.append("%d fila(s) de evol_conc no tienen ESTADO y no se cuentan: %s." % (len(sin_estado), ", ".join(sin_estado)))
    if isinstance(total_excel, (int, float)) and int(total_excel) != len(activas):
        avisos.append("El Excel guardó %d activas (General!C26) y acá hay %d: abrir y guardar el Excel para recalcular."
                      % (int(total_excel), len(activas)))

    mostrar_ab = nombres_unicos(activas, ("Abogado_1", "Abogado_2"))
    mostrar_ec = nombres_unicos(activas, ("Economista_1", "Economista_2"))
    expedientes, sin_ingreso, con_firma = [], [], []
    for f in activas:
        ingreso = a_fecha(f.get("Fecha_Ingreso"))
        if ingreso is None:
            sin_ingreso.append(id_de(f))
            continue
        if a_fecha(f.get("Fecha_firma")):
            con_firma.append(id_de(f))
        expedientes.append({
            "id": id_de(f),
            "operacion": resumir_caratula(f.get("Caratula")),
            "abogados": equipo(f, ("Abogado_1", "Abogado_2"), mostrar_ab),
            "economistas": equipo(f, ("Economista_1", "Economista_2"), mostrar_ec),
            "tipo": TIPO_POR_ES_FT.get(plegar(f.get("ES FT")), "Sin dato"),
            "estado": limpiar(f.get("ESTADO")),
            "ingreso": ingreso.isoformat(),
        })
    if sin_ingreso:
        avisos.append("%d expediente(s) en trámite sin fecha de ingreso, no se muestran: %s." % (len(sin_ingreso), ", ".join(sin_ingreso)))
    if con_firma:
        avisos.append("En trámite pero con fecha de firma cargada (el Excel los cuenta como activos): %s." % ", ".join(con_firma))
    expedientes.sort(key=lambda e: (e["ingreso"], e["id"]))

    return {
        "schema_version": 2,
        "origen": "interno",
        "aviso": AVISO_INTERNO,
        "fuente": fuente,
        "corte": corte.isoformat(),
        "excel_corte": corte_excel.isoformat() if corte_excel else None,
        "estados": list(ESTADOS_ACTIVOS),
        "avisos": avisos,
        "expedientes": expedientes,
    }


def escribir(datos: dict, destino: Path) -> None:
    """Un expediente por linea: los cambios mensuales se ven bien en un diff."""
    cabecera = ["schema_version", "origen", "aviso", "fuente", "corte", "excel_corte", "estados", "avisos"]
    lineas = ["{"]
    for k in cabecera:
        lineas.append("  %s: %s," % (json.dumps(k), json.dumps(datos[k], ensure_ascii=False)))
    lineas.append('  "expedientes": [')
    n = len(datos["expedientes"])
    for i, e in enumerate(datos["expedientes"]):
        lineas.append("    %s%s" % (json.dumps(e, ensure_ascii=False), "," if i < n - 1 else ""))
    lineas += ["  ]", "}", ""]
    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_text("\n".join(lineas), encoding="utf-8")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--excel", default=str(EXCEL), help="Res_firmadas.xlsm (por defecto, junto a este script)")
    ap.add_argument("--salida", default=str(SALIDA), help="JSON de salida")
    ap.add_argument("--corte", default=None, help="fecha para medir la antigüedad, AAAA-MM-DD (por defecto, hoy)")
    args = ap.parse_args()
    corte = a_fecha(args.corte) if args.corte else dt.date.today()
    if corte is None:
        raise SystemExit("--corte tiene que ser AAAA-MM-DD")

    filas, total_excel, corte_excel = leer_excel(Path(args.excel))
    datos = construir(filas, corte, total_excel, corte_excel)
    escribir(datos, Path(args.salida))

    exp = datos["expedientes"]
    print("Expedientes en trámite: %d (el Excel guardó %s) | corte %s | Excel calculado el %s"
          % (len(exp), total_excel, datos["corte"], datos["excel_corte"] or "?"))
    for estado, n in collections.Counter(e["estado"] for e in exp).most_common():
        print("  %-28s %3d" % (estado, n))
    print("  Procedimiento:", dict(collections.Counter(e["tipo"] for e in exp)))
    print("  Abogados: %d | economistas: %d | sin abogado: %d | sin economista: %d" % (
        len({a for e in exp for a in e["abogados"]}), len({a for e in exp for a in e["economistas"]}),
        sum(1 for e in exp if not e["abogados"]), sum(1 for e in exp if not e["economistas"])))
    for a in datos["avisos"]:
        print("  AVISO:", a)
    print("Escrito:", args.salida)
    return 0


if __name__ == "__main__":
    sys.exit(main())
