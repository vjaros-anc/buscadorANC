# -*- coding: utf-8 -*-
"""Genera buscador_opis.html a partir de opi_causas.xlsx.

Uso:
    python 04_html.py                    # usa opi_causas.xlsx -> buscador_opis.html
    python 04_html.py otro.xlsx          # otro Excel de entrada
    python 04_html.py otro.xlsx sal.html # entrada y salida explicitas

Lee la hoja CAUSAS (una fila por OPI) y la hoja NORMAS (normas invocadas por OPI)
y produce un unico archivo HTML autocontenido, sin dependencias externas.
"""
import io
import json
import os
import sys

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
PRE = '<!DOCTYPE html>\n<html lang="es">\n<head>\n<meta charset="utf-8">\n<meta name="viewport" content="width=device-width, initial-scale=1">\n<title>Buscador de OPIs — CNDC</title>\n<style>\n:root{\n  --bg:#f5f6f8; --card:#fff; --ink:#1a1d21; --muted:#666e79; --line:#e2e5ea;\n  --accent:#1f5fa8; --accent-soft:#e8f0fa;\n}\n*{box-sizing:border-box}\nbody{margin:0;font:15px/1.55 -apple-system,Segoe UI,Roboto,Helvetica,Arial,sans-serif;background:var(--bg);color:var(--ink)}\nheader{background:#fff;border-bottom:1px solid var(--line);padding:18px 20px;position:sticky;top:0;z-index:10}\n.wrap{max-width:1100px;margin:0 auto}\nh1{margin:0 0 12px;font-size:20px;font-weight:650}\nh1 span{font-weight:400;color:var(--muted);font-size:14px;margin-left:8px}\n#q{width:100%;padding:11px 14px;font-size:16px;border:1px solid var(--line);border-radius:8px;outline:none}\n#q:focus{border-color:var(--accent);box-shadow:0 0 0 3px var(--accent-soft)}\n.filters{display:flex;flex-wrap:wrap;gap:8px;margin-top:10px}\nselect{padding:7px 10px;border:1px solid var(--line);border-radius:7px;background:#fff;font-size:13px;color:var(--ink);max-width:260px}\n#count{margin-top:8px;font-size:13px;color:var(--muted)}\nmain{max-width:1100px;margin:18px auto;padding:0 20px 60px}\n.card{background:var(--card);border:1px solid var(--line);border-radius:10px;margin-bottom:12px;overflow:hidden}\n.head{padding:14px 16px;cursor:pointer;display:flex;gap:12px;align-items:flex-start}\n.head:hover{background:#fafbfc}\n.num{font-weight:700;color:var(--accent);font-size:15px;min-width:74px}\n.num small{display:block;font-weight:400;color:var(--muted);font-size:11px}\n.summary{flex:1;min-width:0}\n.preg{font-size:14px;margin:0 0 6px}\n.tags{display:flex;flex-wrap:wrap;gap:6px}\n.tag{font-size:11px;padding:2px 8px;border-radius:20px;background:#eef1f5;color:#4a525c;white-space:nowrap}\n.tag.s{background:var(--accent-soft);color:var(--accent);font-weight:600}\n.chev{color:var(--muted);font-size:12px;padding-top:3px}\n.body{display:none;padding:0 16px 16px;border-top:1px solid var(--line)}\n.card.open .body{display:block}\n.f{margin-top:14px}\n.f h3{margin:0 0 4px;font-size:11px;letter-spacing:.06em;text-transform:uppercase;color:var(--muted);font-weight:650}\n.f p{margin:0;font-size:14px;white-space:pre-wrap}\n.f ul{margin:4px 0 0;padding-left:18px;font-size:13.5px}\nmark{background:#ffe9a8;padding:0 1px;border-radius:2px}\n.empty{text-align:center;color:var(--muted);padding:50px 0}\n@media(max-width:640px){.num{min-width:0}.head{flex-wrap:wrap}}\n</style>\n</head>\n<body>\n<header><div class="wrap">\n  <h1>Buscador de OPIs <span id="tot"></span></h1>\n  <input id="q" type="search" placeholder="Buscar en pregunta, respuesta, criterio, normas, expediente…" autocomplete="off">\n  <div class="filters">\n    <select id="f-sent"><option value="">Todos los sentidos de respuesta</option></select>\n    <select id="f-tipo"><option value="">Todos los tipos de pregunta</option></select>\n    <select id="f-crit"><option value="">Todos los criterios</option></select>\n    <select id="f-ley"><option value="">Todas las leyes</option></select>\n  </div>\n  <div id="count"></div>\n</div></header>\n<main id="list"></main>\n<script>\nconst DATA = '
POST = ';\nconst $ = s => document.querySelector(s);\nconst esc = s => (s||\'\').replace(/[&<>]/g,c=>({\'&\':\'&amp;\',\'<\':\'&lt;\',\'>\':\'&gt;\'}[c]));\nconst norm = s => (s||\'\').toLowerCase().normalize(\'NFD\').replace(/[\\u0300-\\u036f]/g,\'\');\nconst split = s => (s||\'\').split(\'|\').map(x=>x.trim()).filter(Boolean);\n\nDATA.forEach(d=>{ d._blob = norm([d.OPI,d.expediente,d.ley,d.pregunta_tipo,d.pregunta_texto,d.argumento_consultante,\n  d.respuesta_sentido,d.respuesta_texto,d.medidas_accesorias,d.CAUSA_ratio_decidendi,d.causa_criterio,\n  d.hecho_determinante,d.fundamento_normativo,d.precedentes,(d.normas_lista||[]).join(\' \')].join(\' \')); });\n\nfunction fill(sel, vals){ vals.sort((a,b)=>a.localeCompare(b,\'es\')).forEach(v=>{\n  const o=document.createElement(\'option\'); o.value=v; o.textContent=v; $(sel).appendChild(o); }); }\nconst uniq = f => [...new Set(DATA.flatMap(f))];\nfill(\'#f-sent\', uniq(d=>split(d.respuesta_sentido)));\nfill(\'#f-tipo\', uniq(d=>split(d.pregunta_tipo)));\nfill(\'#f-crit\', uniq(d=>split(d.causa_criterio)));\nfill(\'#f-ley\',  uniq(d=>split(d.ley)));\n$(\'#tot\').textContent = DATA.length + \' dictámenes\';\n\nfunction hl(txt, terms){\n  let h = esc(txt||\'\');\n  terms.forEach(t=>{ if(t.length<2) return;\n    h = h.replace(new RegExp(\'(\'+t.replace(/[^\\p{L}\\p{N}]/gu, m => \'\\\\\' + m)+\')\',\'gi\'),\'<mark>$1</mark>\'); });\n  return h;\n}\nfunction field(label, val, terms){\n  if(!val) return \'\';\n  return `<div class="f"><h3>${label}</h3><p>${hl(val,terms)}</p></div>`;\n}\nfunction listField(label, arr, terms){\n  if(!arr || !arr.length) return \'\';\n  return `<div class="f"><h3>${label}</h3><ul>${arr.map(x=>`<li>${hl(x,terms)}</li>`).join(\'\')}</ul></div>`;\n}\n\nfunction render(){\n  const raw = $(\'#q\').value.trim();\n  const terms = raw ? raw.split(/\\s+/) : [];\n  const nterms = terms.map(norm);\n  const fs=$(\'#f-sent\').value, ft=$(\'#f-tipo\').value, fc=$(\'#f-crit\').value, fl=$(\'#f-ley\').value;\n  const res = DATA.filter(d =>\n    nterms.every(t => d._blob.includes(t)) &&\n    (!fs || split(d.respuesta_sentido).includes(fs)) &&\n    (!ft || split(d.pregunta_tipo).includes(ft)) &&\n    (!fc || split(d.causa_criterio).includes(fc)) &&\n    (!fl || split(d.ley).includes(fl))\n  );\n  $(\'#count\').textContent = res.length + (res.length===1?\' resultado\':\' resultados\');\n  $(\'#list\').innerHTML = res.length ? res.map(d=>`\n    <div class="card">\n      <div class="head">\n        <div class="num">OPI ${esc(d.OPI)}<small>${esc(d.confianza)}</small></div>\n        <div class="summary">\n          <p class="preg">${hl(d.pregunta_texto,terms)}</p>\n          <div class="tags">\n            ${split(d.respuesta_sentido).map(x=>`<span class="tag s">${esc(x)}</span>`).join(\'\')}\n            ${split(d.pregunta_tipo).map(x=>`<span class="tag">${esc(x)}</span>`).join(\'\')}\n            ${split(d.ley).map(x=>`<span class="tag">Ley ${esc(x)}</span>`).join(\'\')}\n            ${d.expediente?`<span class="tag">${esc(d.expediente)}</span>`:\'\'}\n          </div>\n        </div>\n        <div class="chev">▾</div>\n      </div>\n      <div class="body">\n        ${field(\'Argumento del consultante\', d.argumento_consultante, terms)}\n        ${field(\'Respuesta de la CNDC\', d.respuesta_texto, terms)}\n        ${field(\'Medidas accesorias\', d.medidas_accesorias.replace(/\\|/g,\' · \'), terms)}\n        ${field(\'Ratio decidendi\', d.CAUSA_ratio_decidendi, terms)}\n        ${field(\'Criterio\', d.causa_criterio.replace(/\\|/g,\' · \'), terms)}\n        ${field(\'Hecho determinante\', d.hecho_determinante, terms)}\n        ${field(\'Fundamento normativo\', d.fundamento_normativo.replace(/\\|/g,\' · \'), terms)}\n        ${listField(\'Precedentes citados\', split(d.precedentes), terms)}\n        ${listField(\'Normas invocadas\', d.normas_lista, terms)}\n        ${field(\'Párrafos fuente\', d.parrafos.replace(/\\|/g,\', \'), terms)}\n        ${field(\'Archivo\', d.archivo, terms)}\n      </div>\n    </div>`).join(\'\') : \'<div class="empty">Sin resultados. Probá con otros términos o limpiá los filtros.</div>\';\n}\n$(\'#list\').addEventListener(\'click\', e=>{\n  const h = e.target.closest(\'.head\'); if(h) h.parentElement.classList.toggle(\'open\');\n});\n[\'#q\',\'#f-sent\',\'#f-tipo\',\'#f-crit\',\'#f-ley\'].forEach(s=>$(s).addEventListener(\'input\',render));\nrender();\n</script>\n</body>\n</html>\n'


def leer_causas(wb):
    ws = wb["CAUSAS"]
    filas = list(ws.iter_rows(values_only=True))
    cab = [str(c) for c in filas[0]]
    datos = []
    for fila in filas[1:]:
        if fila[0] is None:
            continue
        datos.append({k: ("" if v is None else str(v).strip())
                      for k, v in zip(cab, fila)})
    return datos


def leer_normas(wb):
    if "NORMAS" not in wb.sheetnames:
        return {}
    normas = {}
    for fila in list(wb["NORMAS"].iter_rows(values_only=True))[1:]:
        if fila[0] is None:
            continue
        opi = str(fila[0]).strip()
        partes = [str(fila[1] or "").strip()]
        if len(fila) > 2 and fila[2]:
            partes.append("art. %s" % str(fila[2]).strip())
        if len(fila) > 3 and fila[3]:
            partes.append("inc. %s" % str(fila[3]).strip())
        texto = " ".join(p for p in partes if p)
        lista = normas.setdefault(opi, [])
        if texto and texto not in lista:
            lista.append(texto)
    return normas


def generar(xlsx, salida):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    datos = leer_causas(wb)
    normas = leer_normas(wb)
    for d in datos:
        d["normas_lista"] = normas.get(d.get("OPI", ""), [])

    # U+2028/2029 son saltos de linea validos en JS y romperian el script
    payload = json.dumps(datos, ensure_ascii=False)
    payload = payload.replace(u"\u2028", "\\u2028")
    payload = payload.replace(u"\u2029", "\\u2029")
    payload = payload.replace("</script", "<\\/script")

    io.open(salida, "w", encoding="utf-8").write(PRE + payload + POST)
    print("%s OPIs -> %s" % (len(datos), salida))


if __name__ == "__main__":
    entrada = sys.argv[1] if len(sys.argv) > 1 else os.path.join(AQUI, "opi_causas.xlsx")
    destino = sys.argv[2] if len(sys.argv) > 2 else os.path.join(AQUI, "buscador_opis.html")
    if not os.path.exists(entrada):
        sys.exit("No se encontro el Excel: %s" % entrada)
    generar(entrada, destino)
